from __future__ import annotations

import hashlib
import json
import re
import uuid
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .models import (
    BUILDER_VERSION,
    DATASET_VERSION,
    REVIEW_SET_NAME,
    SOURCE_VERSION,
)
from .review_decision_loader import DEFAULT_REVIEW_WORKBOOK, load_review_decisions, normalize_cas


DEFAULT_MASTER_WORKBOOK = Path("data/master/msds_index_V24.xlsx")
DEFAULT_GENERATED_DIR = Path("data/generated")
PROFILE_NAMESPACE = uuid.UUID("9152144b-77e0-4f20-a00e-1550452e8fde")
RELATION_NAMESPACE = uuid.UUID("60ef6a46-2d55-44da-8374-d5e752594116")
RULE_NAMESPACE = uuid.UUID("78cbfc6e-8509-4720-9702-438a56861dca")
CAS_PATTERN = re.compile(r"^\d{2,7}-\d{2}-\d$")
SORT_PATTERN = re.compile(r"^(?P<category>\d[A-Z])-(?P<group>\d{2})-(?P<order>\d{3})$")


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _bool_mark(value: Any) -> bool | None:
    raw = _text(value)
    if not raw:
        return None
    return raw in {"○", "Y", "예", "TRUE", "True", "true", "1"}


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _normalize_name(value: Any) -> str:
    return re.sub(r"[\s\-_·,./()\[\]{}]+", "", _text(value)).lower()


def _base_name(name: str) -> str:
    return re.sub(r"\s*[\(（].*$", "", name).strip() or name


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


def _read_master_rows(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[workbook.sheetnames[0]]
        iterator = worksheet.iter_rows(values_only=True)
        headers = [_text(value) for value in next(iterator)]
        rows: list[dict[str, Any]] = []
        for row_number, values in enumerate(iterator, start=2):
            record = {
                header: _json_value(values[index] if index < len(values) else None)
                for index, header in enumerate(headers)
            }
            if any(_text(value) for value in record.values()):
                record["__source_sheet"] = worksheet.title
                record["__source_row_number"] = row_number
                rows.append(record)
        return headers, rows
    finally:
        workbook.close()


class ProfileIdRegistry:
    def __init__(self, path: Path):
        self.path = path
        self.entries: list[dict[str, Any]] = []
        if path.exists():
            data = json.loads(path.read_text(encoding="utf-8"))
            self.entries = list(data.get("entries", []))

    def _find(
        self,
        source_kind: str,
        sort_code: str,
        name: str,
        cas_raw: str,
        review_key: str = "",
    ) -> dict[str, Any] | None:
        for entry in self.entries:
            if review_key and entry.get("review_key") == review_key:
                return entry
        if sort_code:
            matches = [
                entry
                for entry in self.entries
                if entry.get("source_kind") == source_kind
                and entry.get("original_sort_code") == sort_code
            ]
            if len(matches) == 1:
                return matches[0]
            # A V24 sort code is the immutable row key.  Do not fall back to
            # name+CAS when a distinct coded row happens to share both values.
            return None
        normalized_name = _normalize_name(name)
        matches = [
            entry
            for entry in self.entries
            if entry.get("source_kind") == source_kind
            and entry.get("normalized_original_name") == normalized_name
            and entry.get("original_cas_raw", "") == cas_raw
        ]
        return matches[0] if len(matches) == 1 else None

    def assign(
        self,
        *,
        source_kind: str,
        source_version: str,
        source_sheet: str,
        sort_code: str,
        name: str,
        cas_raw: str,
        review_key: str = "",
    ) -> tuple[str, str]:
        existing = self._find(source_kind, sort_code, name, cas_raw, review_key)
        if existing:
            return existing["profile_id"], existing["source_row_id"]

        identity = "|".join(
            [
                source_kind,
                source_version,
                source_sheet,
                sort_code,
                name,
                cas_raw,
                review_key,
            ]
        )
        source_row_id = "SRC_" + uuid.uuid5(PROFILE_NAMESPACE, identity).hex.upper()
        profile_id = "SP_" + uuid.uuid5(PROFILE_NAMESPACE, source_row_id).hex.upper()
        self.entries.append(
            {
                "profile_id": profile_id,
                "source_row_id": source_row_id,
                "source_kind": source_kind,
                "source_version": source_version,
                "source_sheet": source_sheet,
                "original_sort_code": sort_code,
                "original_name": name,
                "normalized_original_name": _normalize_name(name),
                "original_cas_raw": cas_raw,
                "review_key": review_key,
            }
        )
        return profile_id, source_row_id

    def save(self) -> None:
        _write_json(
            self.path,
            {
                "version": 1,
                "description": "Stable profile IDs keyed by immutable source lineage",
                "entries": sorted(self.entries, key=lambda item: item["profile_id"]),
            },
        )


def _decision_state(decision: dict[str, Any] | None, duplicate_cas: bool) -> dict[str, Any]:
    if not decision:
        return {
            "selectable": True,
            "is_parent": False,
            "active": True,
            "default_recommended": False,
            "review_status": "REVIEW_REQUIRED" if duplicate_cas else "AUTO_MIGRATED",
            "user_review_note": "",
            "review_decision": None,
        }
    item_decision = decision["item_decision"]
    mapping = {
        "선택가능": (True, False, True, "APPROVED"),
        # Parent/representative records remain active as detection/grouping
        # anchors, but are not an extra user-selectable measurement factor.
        "부모·대표": (False, True, True, "APPROVED"),
        "별칭·중복": (False, False, True, "APPROVED"),
        "제외": (False, False, False, "EXCLUDED"),
        "보류": (False, False, False, "REVIEW_REQUIRED"),
    }
    selectable, is_parent, active, status = mapping.get(
        item_decision, (False, False, False, "REVIEW_REQUIRED")
    )
    return {
        "selectable": selectable,
        "is_parent": is_parent,
        "active": active,
        "default_recommended": decision["default_recommended"],
        "review_status": status,
        "user_review_note": decision["user_review_note"],
        "review_decision": item_decision,
    }


def _sort_fields(sort_code: str, canonical_name: str, profile_id: str) -> dict[str, Any]:
    match = SORT_PATTERN.match(sort_code)
    if match:
        category_order = int(match.group("category")[0])
        sort_group = f"{match.group('category')}-{match.group('group')}"
        sort_order = int(match.group("order"))
    else:
        category_order, sort_group, sort_order = 999, sort_code or "ZZZ", 999999
    return {
        "category_order": category_order,
        "sort_group": sort_group,
        "sort_order": sort_order,
        "sort_key": [category_order, sort_group, sort_order, canonical_name, profile_id],
    }


def _profile_from_master(
    row: dict[str, Any],
    decision: dict[str, Any] | None,
    duplicate_cas: bool,
    registry: ProfileIdRegistry,
) -> dict[str, Any]:
    original_name = _text(row.get("측정대상 물질명"))
    canonical_name = decision["canonical_name"] if decision else original_name
    cas_raw = _text(row.get("CAS No."))
    cas_no = normalize_cas(cas_raw)
    sort_code = _text(row.get("정렬코드"))
    profile_id, source_row_id = registry.assign(
        source_kind="master",
        source_version=SOURCE_VERSION,
        source_sheet=row["__source_sheet"],
        sort_code=sort_code,
        name=original_name,
        cas_raw=cas_raw,
    )
    state = _decision_state(decision, duplicate_cas)
    profile = {
        "profile_id": profile_id,
        "profile_role": "measurement_factor",
        "source_sheet": row["__source_sheet"],
        "source_row_number": row["__source_row_number"],
        "source_row_id": source_row_id,
        "source_sort_code": sort_code,
        "cas_no": cas_no,
        "cas_no_raw": cas_raw,
        "canonical_name": canonical_name,
        "source_name_raw": original_name,
        "base_name": _base_name(canonical_name),
        "parent_profile_id": None,
        **state,
        "category_code": _text(row.get("대분류")),
        "analysis_method": _text(row.get("분석방법")),
        "sampling_media": _text(row.get("매체(실무용)")),
        "sampling_media_reference": _text(row.get("매체")),
        "preferred_flow_rate": _text(row.get("적정유속(L/min)")),
        "flow_rate": _text(row.get("유속(L/min)")),
        "minimum_air_volume_l_raw": _text(row.get("최소공기량(L)")),
        "maximum_air_volume_l_raw": _text(row.get("최대공기량(L)")),
        "sample_count_raw": _text(row.get("시료수")),
        "twa": _text(row.get("노출기준(TWA)")),
        "stel": _text(row.get("노출기준(STEL)")),
        "management_target": _bool_mark(row.get("측정")),
        "special_management": _bool_mark(row.get("특별관리")),
        "permit_required": _bool_mark(row.get("허가대상")),
        "health_exam_target": _bool_mark(row.get("특검")),
        "carcinogenicity": _text(row.get("발암성")),
        "germ_cell_mutagenicity": _text(row.get("생식세포")),
        "reproductive_toxicity": _text(row.get("생식독성")),
        "analysis_available": _text(row.get("분석가능여부")),
        "analysis_notes": _text(row.get("분석관련 특이사항")),
        "notes": _text(row.get("비고")),
        "source_version": SOURCE_VERSION,
        "source_review_set": REVIEW_SET_NAME if decision else None,
        "source_raw": {
            key: value for key, value in row.items() if not key.startswith("__")
        },
    }
    profile.update(_sort_fields(sort_code, canonical_name, profile_id))
    return profile


def _find_profile(
    profiles: list[dict[str, Any]], name: str, cas_no: str | None
) -> dict[str, Any] | None:
    candidates = [
        profile
        for profile in profiles
        if (not cas_no or profile.get("cas_no") == cas_no)
    ]
    wanted = _normalize_name(name)
    exact = [
        profile
        for profile in candidates
        if _normalize_name(profile.get("canonical_name")) == wanted
    ]
    if len(exact) == 1:
        return exact[0]
    contained = [
        profile
        for profile in candidates
        if _normalize_name(profile.get("canonical_name"))
        and _normalize_name(profile.get("canonical_name")) in wanted
    ]
    contained.sort(key=lambda item: len(_normalize_name(item["canonical_name"])), reverse=True)
    return contained[0] if contained else (candidates[0] if len(candidates) == 1 else None)


def _supplemental_source_profile(
    decision: dict[str, Any], registry: ProfileIdRegistry
) -> dict[str, Any]:
    review_key = f"{decision['candidate_id']}:source:{decision['source_cas'] or decision['source_name']}"
    profile_id, source_row_id = registry.assign(
        source_kind="review_detection_source",
        source_version=REVIEW_SET_NAME,
        source_sheet=decision["review_sheet"],
        sort_code="",
        name=decision["source_name"],
        cas_raw=decision["source_cas_raw"],
        review_key=review_key,
    )
    active = decision["final_decision"] in {"승인", "수정"}
    profile = {
        "profile_id": profile_id,
        "profile_role": "detection_source",
        "source_sheet": decision["review_sheet"],
        "source_row_number": decision["review_row_number"],
        "source_row_id": source_row_id,
        "source_sort_code": "",
        "cas_no": decision["source_cas"],
        "cas_no_raw": decision["source_cas_raw"],
        "canonical_name": decision["source_name"],
        "source_name_raw": decision["source_name"],
        "base_name": _base_name(decision["source_name"]),
        "parent_profile_id": None,
        "selectable": False,
        "is_parent": False,
        "active": active,
        "default_recommended": False,
        "category_code": "DETECTION_SOURCE",
        "analysis_method": "",
        "sampling_media": "",
        "twa": "",
        "stel": "",
        "management_target": None,
        "special_management": None,
        "permit_required": None,
        "health_exam_target": None,
        "review_status": "APPROVED" if active else "REVIEW_REQUIRED",
        "review_decision": decision["final_decision"],
        "user_review_note": decision["user_review_note"],
        "source_version": REVIEW_SET_NAME,
        "source_review_set": REVIEW_SET_NAME,
        "source_raw": decision["source_record"],
    }
    profile.update(_sort_fields("", profile["canonical_name"], profile_id))
    return profile


def _relation_id(source_id: str, target_id: str, relation_type: str, key: str) -> str:
    return "REL_" + uuid.uuid5(
        RELATION_NAMESPACE, f"{source_id}|{target_id}|{relation_type}|{key}"
    ).hex.upper()


def _rule_id(candidate_id: str) -> str:
    return "RULE_" + uuid.uuid5(RULE_NAMESPACE, candidate_id).hex.upper()


def _selection_source(relation_type: str) -> str:
    if relation_type == "ANALYTICAL_RELATION" or relation_type == "MANUAL_REVIEW":
        return "analytical_relation"
    return "process_recommended"


def _build_same_cas_relations(
    profiles: list[dict[str, Any]], reviewed_sort_codes: set[str]
) -> list[dict[str, Any]]:
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for profile in profiles:
        if profile["source_sort_code"] in reviewed_sort_codes and profile.get("cas_no"):
            groups[profile["cas_no"]].append(profile)

    relations: list[dict[str, Any]] = []
    for cas_no, group in groups.items():
        active_group = [profile for profile in group if profile["active"]]
        active_group.sort(key=lambda item: item["sort_key"])
        if not active_group:
            continue
        anchor = active_group[0]
        for target in active_group[1:]:
            relations.append(
                {
                    "relation_id": _relation_id(
                        anchor["profile_id"], target["profile_id"], "SAME_CAS_VARIANT", cas_no
                    ),
                    "source_profile_id": anchor["profile_id"],
                    "target_profile_id": target["profile_id"],
                    "relation_type": "SAME_CAS_VARIANT",
                    "bidirectional": True,
                    "allow_source_only": True,
                    "allow_target_only": True,
                    "allow_multiple": True,
                    "default_selected": bool(target["default_recommended"]),
                    "require_user_confirmation": True,
                    "active": True,
                    "review_status": "APPROVED",
                    "guide_message": f"동일 CAS {cas_no}의 별도 측정인자 후보",
                    "user_review_note": target["user_review_note"],
                    "source_review_set": REVIEW_SET_NAME,
                }
            )

        parents = [profile for profile in active_group if profile["is_parent"]]
        if len(parents) == 1:
            parent = parents[0]
            for child in active_group:
                if child["profile_id"] == parent["profile_id"] or child["is_parent"]:
                    continue
                child["parent_profile_id"] = parent["profile_id"]
                relations.append(
                    {
                        "relation_id": _relation_id(
                            parent["profile_id"], child["profile_id"], "PARENT_CHILD", cas_no
                        ),
                        "source_profile_id": parent["profile_id"],
                        "target_profile_id": child["profile_id"],
                        "relation_type": "PARENT_CHILD",
                        "bidirectional": False,
                        "allow_source_only": True,
                        "allow_target_only": True,
                        "allow_multiple": True,
                        "default_selected": bool(child["default_recommended"]),
                        "require_user_confirmation": True,
                        "active": True,
                        "review_status": "APPROVED",
                        "guide_message": f"{parent['canonical_name']} 그룹의 세부 측정인자",
                        "user_review_note": child["user_review_note"],
                        "source_review_set": REVIEW_SET_NAME,
                    }
                )
    return relations


def build_dataset(
    repo_root: str | Path,
    *,
    master_path: str | Path = DEFAULT_MASTER_WORKBOOK,
    review_path: str | Path = DEFAULT_REVIEW_WORKBOOK,
    output_dir: str | Path = DEFAULT_GENERATED_DIR,
) -> dict[str, Any]:
    root = Path(repo_root).resolve()
    master = (root / master_path).resolve() if not Path(master_path).is_absolute() else Path(master_path)
    review = (root / review_path).resolve() if not Path(review_path).is_absolute() else Path(review_path)
    generated = (root / output_dir).resolve() if not Path(output_dir).is_absolute() else Path(output_dir)
    if not master.exists():
        raise FileNotFoundError(f"원본 마스터가 없습니다: {master}")
    if not review.exists():
        raise FileNotFoundError(f"승인 파일이 없습니다: {review}")

    generated.mkdir(parents=True, exist_ok=True)
    review_result = load_review_decisions(review)
    _, master_rows = _read_master_rows(master)
    if len(master_rows) != 297:
        raise ValueError(f"V24 데이터 행 수가 297이 아님: {len(master_rows)}")

    decisions_by_sort = {
        decision["source_sort_code"]: decision
        for decision in review_result.same_cas_decisions
    }
    cas_counts = Counter(
        normalize_cas(row.get("CAS No."))
        for row in master_rows
        if normalize_cas(row.get("CAS No."))
    )
    registry = ProfileIdRegistry(generated / "profile_id_registry.json")
    profiles = [
        _profile_from_master(
            row,
            decisions_by_sort.get(_text(row.get("정렬코드"))),
            bool(normalize_cas(row.get("CAS No.")) and cas_counts[normalize_cas(row.get("CAS No."))] > 1),
            registry,
        )
        for row in master_rows
    ]

    for decision in review_result.cross_cas_decisions:
        source = _find_profile(profiles, decision["source_name"], decision["source_cas"])
        if source is None and decision["final_decision"] != "제외":
            supplemental = _supplemental_source_profile(decision, registry)
            existing = _find_profile(
                profiles, supplemental["canonical_name"], supplemental["cas_no"]
            )
            if existing is None:
                profiles.append(supplemental)

    same_relations = _build_same_cas_relations(
        profiles, set(decisions_by_sort)
    )
    relations = list(same_relations)
    rules: list[dict[str, Any]] = []
    conversion_warnings = [warning.to_dict() for warning in review_result.warnings]

    for decision in review_result.cross_cas_decisions:
        source = _find_profile(profiles, decision["source_name"], decision["source_cas"])
        target = _find_profile(profiles, decision["target_name"], decision["target_cas"])
        final_decision = decision["final_decision"]
        policies_complete = (
            decision["default_selected"] is not None
            and decision["allow_multiple"] is not None
        )
        approved = final_decision in {"승인", "수정"} and policies_complete
        review_status = (
            "EXCLUDED"
            if final_decision == "제외"
            else "APPROVED"
            if approved
            else "REVIEW_REQUIRED"
        )
        if source is None or target is None:
            approved = False
            review_status = "REVIEW_REQUIRED" if final_decision != "제외" else "EXCLUDED"
            conversion_warnings.append(
                {
                    "code": "PROFILE_DIRECTION_UNRESOLVED",
                    "message": f"{decision['candidate_id']} source/target profile을 확정하지 못함",
                    "sheet": decision["review_sheet"],
                    "row": decision["review_row_number"],
                    "field": None,
                }
            )
        source_id = source["profile_id"] if source else None
        target_id = target["profile_id"] if target else None
        relation_key = decision["candidate_id"]
        relations.append(
            {
                "relation_id": _relation_id(
                    source_id or "UNRESOLVED",
                    target_id or "UNRESOLVED",
                    decision["relation_type"],
                    relation_key,
                ),
                "source_profile_id": source_id,
                "target_profile_id": target_id,
                "relation_type": decision["relation_type"],
                "bidirectional": False,
                "allow_source_only": True,
                "allow_target_only": True,
                "allow_multiple": decision["allow_multiple"],
                "default_selected": decision["default_selected"],
                "require_user_confirmation": bool(
                    decision["require_user_confirmation"] is not False
                ),
                "active": approved,
                "review_status": review_status,
                "guide_message": decision["evidence_note"],
                "user_review_note": decision["user_review_note"],
                "source_review_set": REVIEW_SET_NAME,
                "source_review_id": decision["candidate_id"],
                "final_decision": final_decision,
            }
        )
        if final_decision == "제외":
            continue
        rules.append(
            {
                "rule_id": _rule_id(decision["candidate_id"]),
                "source_profile_id": source_id,
                "source_cas": decision["source_cas"],
                "source_name": decision["source_name"],
                "target_profile_id": target_id,
                "target_cas": target.get("cas_no") if target else decision["target_cas"],
                "target_name": target.get("canonical_name") if target else decision["target_name"],
                "relation_type": decision["relation_type"],
                "trigger_condition": decision["trigger_condition"],
                "condition_or_level_amendment": decision[
                    "condition_or_level_amendment"
                ],
                "recommendation_level": decision["recommendation_level"],
                "default_selected": decision["default_selected"],
                "allow_multiple": decision["allow_multiple"],
                "require_user_confirmation": bool(
                    decision["require_user_confirmation"] is not False
                ),
                "selection_source": _selection_source(decision["relation_type"]),
                "active": approved,
                "review_status": review_status,
                "auto_confirmed": False,
                "user_review_note": decision["user_review_note"],
                "evidence_note": decision["evidence_note"],
                "source_review_set": REVIEW_SET_NAME,
                "source_review_id": decision["candidate_id"],
                "final_decision": final_decision,
            }
        )

    profiles.sort(key=lambda item: item["sort_key"])
    relations.sort(key=lambda item: item["relation_id"])
    rules.sort(key=lambda item: item["rule_id"])
    registry.save()

    from .index_validator import validate_dataset

    report = validate_dataset(
        profiles,
        relations,
        rules,
        repo_root=root,
        source_master_rows=len(master_rows),
        expected_source_master_rows=297,
        loaded_review_decisions=review_result.decision_count,
        converted_review_decisions=len(review_result.same_cas_decisions)
        + len(review_result.cross_cas_decisions),
        reference_paths=[master, review],
        check_old_paths=True,
    )

    _write_json(generated / "substance_profiles.json", profiles)
    _write_json(generated / "substance_relations.json", relations)
    _write_json(generated / "recommendation_rules.json", rules)
    manifest = {
        "dataset_version": DATASET_VERSION,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source_master_path": master.relative_to(root).as_posix(),
        "source_master_sha256": _sha256(master),
        "review_set_path": (review.parent.relative_to(root)).as_posix(),
        "review_workbook_path": review.relative_to(root).as_posix(),
        "review_workbook_sha256": _sha256(review),
        "source_master_row_count": len(master_rows),
        "staging_row_count": len(master_rows),
        "profile_count": len(profiles),
        "master_profile_count": sum(
            profile["source_version"] == SOURCE_VERSION for profile in profiles
        ),
        "supplemental_detection_profile_count": sum(
            profile["profile_role"] == "detection_source" for profile in profiles
        ),
        "relation_count": len(relations),
        "active_recommendation_rule_count": sum(rule["active"] for rule in rules),
        "held_rule_count": sum(
            rule["final_decision"] == "보류" for rule in rules
        ),
        "excluded_rule_count": sum(
            decision["final_decision"] == "제외"
            for decision in review_result.cross_cas_decisions
        ),
        "incomplete_policy_rule_count": sum(
            rule["review_status"] == "REVIEW_REQUIRED"
            and rule["final_decision"] in {"승인", "수정"}
            for rule in rules
        ),
        "same_cas_review_decision_count": len(review_result.same_cas_decisions),
        "cross_cas_review_decision_count": len(review_result.cross_cas_decisions),
        "review_loader_warnings": conversion_warnings,
        "validator": report.to_dict(),
        "builder_version": BUILDER_VERSION,
    }
    _write_json(generated / "dataset_manifest.json", manifest)
    return {
        "profiles": profiles,
        "relations": relations,
        "rules": rules,
        "manifest": manifest,
        "validation_report": report,
    }


def main() -> int:
    root = Path(__file__).resolve().parents[1]
    result = build_dataset(root)
    report = result["validation_report"]
    print(
        json.dumps(
            {
                "profiles": len(result["profiles"]),
                "relations": len(result["relations"]),
                "rules": len(result["rules"]),
                "active_rules": sum(rule["active"] for rule in result["rules"]),
                "errors": len(report.errors),
                "warnings": len(report.warnings),
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0 if report.is_valid else 1


if __name__ == "__main__":
    raise SystemExit(main())
