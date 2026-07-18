from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from .models import REVIEW_SET_NAME, ReviewLoadResult, ReviewWarning


DEFAULT_REVIEW_WORKBOOK = Path(
    "data/review"
) / REVIEW_SET_NAME / "MSDS_측정인자_1차검토표_최종수정.xlsx"

SAME_CAS_SHEET = "동일CAS_우선검토"
CROSS_CAS_SHEET = "교차CAS_16건"
CAS_PATTERN = re.compile(r"(?<!\d)(\d{2,7}-\d{2}-\d)(?!\d)")

SAME_CAS_REQUIRED = {
    "검토그룹",
    "CAS",
    "정렬코드",
    "측정대상 물질명",
    "항목판정",
    "기본추천",
    "사용자 의견",
}
CROSS_CAS_REQUIRED = {
    "번호",
    "원인자",
    "원CAS",
    "추천인자",
    "추천CAS",
    "최종결정",
    "수정 조건/추천강도",
    "기본체크",
    "복수선택",
    "사용자 의견",
}


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def normalize_cas(value: Any) -> str | None:
    """Return a standard CAS search key while retaining raw values elsewhere."""
    raw = _text(value)
    if not raw:
        return None
    match = CAS_PATTERN.search(raw.replace(" ", ""))
    return match.group(1) if match else None


def normalize_yes_no(value: Any) -> bool | None:
    raw = _text(value).upper()
    if raw == "Y":
        return True
    if raw == "N":
        return False
    if raw in {"", "해당없음"}:
        return None
    return None


def _records(worksheet: Any) -> Iterable[tuple[int, dict[str, Any]]]:
    rows = worksheet.iter_rows(values_only=True)
    try:
        header_values = next(rows)
    except StopIteration:
        return
    headers = [_text(value) for value in header_values]
    for row_number, values in enumerate(rows, start=2):
        record = {
            header: values[index] if index < len(values) else None
            for index, header in enumerate(headers)
            if header
        }
        if any(_text(value) for value in record.values()):
            yield row_number, record


def _validate_headers(
    worksheet: Any, required: set[str], sheet_name: str
) -> None:
    headers = {_text(cell.value) for cell in worksheet[1] if _text(cell.value)}
    missing = sorted(required - headers)
    if missing:
        raise ValueError(f"{sheet_name} 필수 열 누락: {', '.join(missing)}")


def _apply_target_name_amendment(record: dict[str, Any]) -> str:
    target_name = _text(record.get("추천인자"))
    if _text(record.get("최종결정")) != "수정":
        return target_name
    user_note = _text(record.get("사용자 의견"))
    match = re.search(r"추천인자\s*[:：]\s*(.+)$", user_note)
    return match.group(1).strip() if match else target_name


def load_review_decisions(
    workbook_path: str | Path = DEFAULT_REVIEW_WORKBOOK,
) -> ReviewLoadResult:
    path = Path(workbook_path)
    if not path.exists():
        raise FileNotFoundError(f"사용자 승인 파일이 없습니다: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        missing_sheets = {
            SAME_CAS_SHEET,
            CROSS_CAS_SHEET,
        } - set(workbook.sheetnames)
        if missing_sheets:
            raise ValueError(
                "승인 파일 필수 시트 누락: " + ", ".join(sorted(missing_sheets))
            )

        same_sheet = workbook[SAME_CAS_SHEET]
        cross_sheet = workbook[CROSS_CAS_SHEET]
        _validate_headers(same_sheet, SAME_CAS_REQUIRED, SAME_CAS_SHEET)
        _validate_headers(cross_sheet, CROSS_CAS_REQUIRED, CROSS_CAS_SHEET)

        result = ReviewLoadResult(workbook_path=path)

        valid_item_decisions = {"선택가능", "부모·대표", "별칭·중복", "제외", "보류"}
        for row_number, record in _records(same_sheet):
            item_decision = _text(record.get("항목판정"))
            default_raw = _text(record.get("기본추천"))
            cas_raw = _text(record.get("CAS"))
            decision = {
                "review_sheet": SAME_CAS_SHEET,
                "review_row_number": row_number,
                "review_group": _text(record.get("검토그룹")),
                "cas_no": normalize_cas(cas_raw),
                "cas_no_raw": cas_raw,
                "source_sort_code": _text(record.get("정렬코드")),
                "canonical_name": _text(record.get("측정대상 물질명")),
                "item_decision": item_decision,
                "default_recommended": normalize_yes_no(default_raw),
                "default_recommended_raw": default_raw,
                "user_review_note": _text(record.get("사용자 의견")),
                "source_record": {key: value for key, value in record.items()},
            }
            result.same_cas_decisions.append(decision)

            if item_decision not in valid_item_decisions:
                result.warnings.append(
                    ReviewWarning(
                        "UNKNOWN_ITEM_DECISION",
                        f"알 수 없는 항목판정: {item_decision or '(공란)'}",
                        SAME_CAS_SHEET,
                        row_number,
                        "항목판정",
                    )
                )
            if default_raw not in {"Y", "N", "해당없음"}:
                result.warnings.append(
                    ReviewWarning(
                        "UNKNOWN_DEFAULT_RECOMMENDATION",
                        f"기본추천 값을 해석할 수 없음: {default_raw or '(공란)'}",
                        SAME_CAS_SHEET,
                        row_number,
                        "기본추천",
                    )
                )
            if cas_raw and not decision["cas_no"]:
                result.warnings.append(
                    ReviewWarning(
                        "NONSTANDARD_CAS",
                        f"표준 CAS를 추출할 수 없음: {cas_raw}",
                        SAME_CAS_SHEET,
                        row_number,
                        "CAS",
                    )
                )

        valid_final_decisions = {"승인", "수정", "제외", "보류"}
        for row_number, record in _records(cross_sheet):
            final_decision = _text(record.get("최종결정"))
            default_raw = _text(record.get("기본체크"))
            multiple_raw = _text(record.get("복수선택"))
            source_cas_raw = _text(record.get("원CAS"))
            target_cas_raw = _text(record.get("추천CAS"))
            target_name = _apply_target_name_amendment(record)
            decision = {
                "review_sheet": CROSS_CAS_SHEET,
                "review_row_number": row_number,
                "candidate_id": _text(record.get("번호")),
                "source_name": _text(record.get("원인자")),
                "source_cas": normalize_cas(source_cas_raw),
                "source_cas_raw": source_cas_raw,
                "target_name": target_name,
                "target_name_raw": _text(record.get("추천인자")),
                "target_cas": normalize_cas(target_cas_raw),
                "target_cas_raw": target_cas_raw,
                "relation_type": _text(record.get("제안 관계")),
                "trigger_condition": _text(record.get("발동 조건")),
                "recommendation_level": _text(record.get("추천강도")),
                "proposed_default_selected": normalize_yes_no(
                    record.get("기본체크(제안)")
                ),
                "proposed_allow_multiple": normalize_yes_no(
                    record.get("복수선택(제안)")
                ),
                "proposed_require_user_confirmation": normalize_yes_no(
                    record.get("사용자확인(제안)")
                ),
                "evidence_note": _text(record.get("근거·현재 동작")),
                "final_decision": final_decision,
                "condition_or_level_amendment": _text(
                    record.get("수정 조건/추천강도")
                ),
                "default_selected": normalize_yes_no(default_raw),
                "default_selected_raw": default_raw,
                "allow_multiple": normalize_yes_no(multiple_raw),
                "allow_multiple_raw": multiple_raw,
                "require_user_confirmation": normalize_yes_no(
                    record.get("사용자확인(제안)")
                ),
                "user_review_note": _text(record.get("사용자 의견")),
                "source_record": {key: value for key, value in record.items()},
            }
            result.cross_cas_decisions.append(decision)

            if final_decision not in valid_final_decisions:
                result.warnings.append(
                    ReviewWarning(
                        "UNKNOWN_FINAL_DECISION",
                        f"알 수 없는 최종결정: {final_decision or '(공란)'}",
                        CROSS_CAS_SHEET,
                        row_number,
                        "최종결정",
                    )
                )
            if final_decision in {"승인", "수정"}:
                for field, raw_value in (
                    ("기본체크", default_raw),
                    ("복수선택", multiple_raw),
                ):
                    if raw_value not in {"Y", "N"}:
                        result.warnings.append(
                            ReviewWarning(
                                "BLANK_RUNTIME_POLICY",
                                f"{field}이 공란이므로 운영 규칙을 활성화하지 않음",
                                CROSS_CAS_SHEET,
                                row_number,
                                field,
                            )
                        )
            for field, raw_value, normalized in (
                ("원CAS", source_cas_raw, decision["source_cas"]),
                ("추천CAS", target_cas_raw, decision["target_cas"]),
            ):
                if raw_value and not normalized:
                    result.warnings.append(
                        ReviewWarning(
                            "NONSTANDARD_CAS",
                            f"표준 CAS를 추출할 수 없음: {raw_value}",
                            CROSS_CAS_SHEET,
                            row_number,
                            field,
                        )
                    )
            if (
                decision["source_cas"]
                and decision["source_cas"] == decision["target_cas"]
                and decision["source_name"] == decision["target_name"]
            ):
                result.warnings.append(
                    ReviewWarning(
                        "SELF_RELATION",
                        "원인자와 추천인자가 동일함",
                        CROSS_CAS_SHEET,
                        row_number,
                    )
                )
        return result
    finally:
        workbook.close()
