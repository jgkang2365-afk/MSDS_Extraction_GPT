import pandas as pd
import json
import os
import re
from collections import defaultdict

def merge_master_datasets():
    excel_path = "msds_index.xlsx"
    json_path = "MES_MASTER_LOOKUP.json"
    output_path = "msds_index_merged.xlsx"

    print(f"[*] 데이터 병합 작업을 시작합니다...")

    # 1. JSON 파일 로드
    try:
        with open(json_path, "r", encoding="utf-8") as f:
            json_data = json.load(f)
    except Exception as e:
        print(f"[!] JSON 파일을 읽는 데 실패했습니다: {e}")
        return

    # CAS 번호 -> 레코드(들) 리스트 매핑 (1:N 대응용)
    lookup_by_cas = defaultdict(list)
    # 이름 정규화 키 -> 레코드 매핑
    lookup_by_name = {}

    def clean_name_key(name_str):
        if not name_str:
            return ""
        return re.sub(r"\s+", "", str(name_str)).lower()

    # JSON 데이터에서 매핑 정보 추출 및 구축
    # 루트 레벨의 CAS 키 탐색
    for key, value in json_data.items():
        if isinstance(value, dict) and "CAS번호" in value:
            cas = str(value["CAS번호"]).strip()
            if cas:
                lookup_by_cas[cas].append(value)
            
            std_name = value.get("물질명")
            if std_name:
                lookup_by_name[clean_name_key(std_name)] = value
            s_name = value.get("상용명")
            if s_name:
                lookup_by_name[clean_name_key(s_name)] = value

    # master_list 배열 내부 탐색
    if "master_list" in json_data:
        for item in json_data["master_list"]:
            cas = str(item.get("CAS번호", "")).strip()
            if cas:
                lookup_by_cas[cas].append(item)
            
            std_name = item.get("물질명")
            if std_name:
                lookup_by_name[clean_name_key(std_name)] = item
            s_name = item.get("상용명")
            if s_name:
                lookup_by_name[clean_name_key(s_name)] = item

    print(f"[*] JSON에서 고유 CAS 번호 수 {len(lookup_by_cas)}개, 이름 매핑 사전 {len(lookup_by_name)}개를 구축했습니다.")

    # 2. 엑셀 파일 로드
    try:
        df = pd.read_excel(excel_path)
        print(f"[*] 엑셀 파일 로드 완료: 총 {len(df)}행")
    except Exception as e:
        print(f"[!] 엑셀 파일을 읽는 데 실패했습니다: {e}")
        return

    # 3. 이관할 JSON 핵심 컬럼 정의
    target_columns = [
        "노출기준(TWA)", "노출기준(STEL)", "측정", "특검", 
        "특별관리", "허가대상", "발암성", "생식세포", "생식독성", "LOD"
    ]

    # 엑셀에 해당 컬럼이 없으면 빈 열로 생성하고, 기존 열이 있다면 모두 초기화
    for col in target_columns:
        df[col] = ""

    # CAS 번호가 없는 물질들에 대한 대체 명칭 매핑 규칙
    FALLBACK_NAME_MAP = {
        "기타분진": "기타분진(유리규산 1%이하)",
        "나무분진-연목": "목재분진(적삼목외 모든종)",
        "나무분진-강목": "목재분진(적삼목외 모든종)",
        "목재분진": "목재분진(적삼목외 모든종)",
        "목재분진(적삼목외 기타 모든 종)": "목재분진(적삼목외 모든종)",
        "6가크롬": "6가크롬(수용성)",
    }

    # 4. 복합 병합(VLOOKUP) 수행
    match_count = 0
    skip_count = 0
    highlight_rows = [] # 유사/대체 매칭이 적용된 행 인덱스 저장 (0-based)
    
    for idx, row in df.iterrows():
        cas_no = str(row.get("CAS No.", "")).strip()
        m_name = str(row.get("측정대상 물질명", "")).strip()

        # 물리적 인자인 소음, 고열은 매칭 및 TWA 정보 입력에서 완전히 배제
        if m_name in ["소음", "고열"]:
            skip_count += 1
            continue

        json_row = None
        is_highlight = False

        # 1순위: 이름 정밀 일치 매칭 (Exact Name Match)
        # 대체 룰 확인
        mapped_name = FALLBACK_NAME_MAP.get(m_name, m_name)
        clean_mapped_name = clean_name_key(mapped_name)
        
        if clean_mapped_name in lookup_by_name:
            json_row = lookup_by_name[clean_mapped_name]
            # 만약 FALLBACK_NAME_MAP을 거쳤다면 수동 검증 대상(하이라이트)으로 분류
            if m_name in FALLBACK_NAME_MAP:
                is_highlight = True
                print(f"[!] 대체 규칙 매칭: {m_name} -> {mapped_name} (노란색 하이라이트)")

        # 2순위: 1순위 실패 시, CAS 번호 및 이름 복합 분석
        if json_row is None:
            has_cas = pd.notna(row.get("CAS No.")) and cas_no != "" and cas_no.lower() != "nan"
            if has_cas:
                cas_clean = re.sub(r"\s+", "", cas_no)
                cas_match = re.search(r"(\d{2,7}-\d{2}-\d)", cas_clean)
                if cas_match:
                    target_cas = cas_match.group(1)
                    candidates = lookup_by_cas.get(target_cas, [])
                    
                    if len(candidates) == 1:
                        # 고유한 CAS 매칭 (중복 없음)
                        json_row = candidates[0]
                    elif len(candidates) > 1:
                        # 1:N 중복 CAS 후보들 발생 (예: 알루미늄 계열 등) -> 이름 유사 분석
                        clean_m_name = clean_name_key(m_name)
                        best_candidate = None
                        
                        # 후보들 중 이름이 부분 일치하는 레코드 탐색
                        for cand in candidates:
                            cand_std = clean_name_key(cand.get("물질명", ""))
                            cand_s = clean_name_key(cand.get("상용명", ""))
                            
                            if clean_m_name in cand_std or cand_std in clean_m_name or clean_m_name in cand_s or cand_s in clean_m_name:
                                best_candidate = cand
                                break
                        
                        if best_candidate:
                            json_row = best_candidate
                            # 복합 매칭(중복 CAS 내 구별) 또한 사용자가 정합성을 체크할 수 있도록 하이라이트
                            is_highlight = True
                            print(f"[!] 중복 CAS 내 성상 필터 매칭 성공: {m_name} (CAS: {target_cas}) -> {json_row.get('물질명')} (노란색 하이라이트)")
                        else:
                            # 부분 매칭도 실패할 경우, 후보군의 첫 번째를 대입하고 강력히 검증 대상으로 분류
                            json_row = candidates[0]
                            is_highlight = True
                            print(f"[?] 중복 CAS 성상 필터 실패: {m_name} (CAS: {target_cas}) -> 기본값 {json_row.get('물질명')} 적용 (노란색 하이라이트)")

        # 데이터 이식
        if json_row:
            match_count += 1
            if is_highlight:
                highlight_rows.append(idx)
                
            for col in target_columns:
                val = json_row.get(col, "")
                if pd.notna(val) and str(val).strip() != "":
                    df.at[idx, col] = str(val).strip()
        else:
            print(f"[?] 매칭 실패 물질: {m_name} (CAS: {cas_no})")

    # 5. 결과 저장 및 스타일링 적용 (openpyxl 사용)
    try:
        # 1차로 pandas 데이터 저장
        df.to_excel(output_path, index=False)
        
        # openpyxl로 열어서 바탕색 스타일링
        from openpyxl import load_workbook
        from openpyxl.styles import PatternFill

        wb = load_workbook(output_path)
        ws = wb.active

        # 연한 노란색 채우기 색상 정의 (Light Yellow)
        yellow_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

        # target_columns 중 "노출기준(TWA)" 컬럼 위치 찾기 (1-based index)
        twa_col_idx = None
        for col_idx in range(1, ws.max_column + 1):
            if ws.cell(row=1, column=col_idx).value == "노출기준(TWA)":
                twa_col_idx = col_idx
                break

        # 하이라이트 대상인 행들의 TWA 셀에 바탕색 칠하기
        # openpyxl의 행은 1-based이고, 첫 번째 줄은 헤더이므로 row = idx + 2 가 됨
        if twa_col_idx:
            for r_idx in highlight_rows:
                cell = ws.cell(row=r_idx + 2, column=twa_col_idx)
                cell.fill = yellow_fill

        wb.save(output_path)
        print(f"\n[+] 병합 완료! 총 {match_count}개의 행에 규제 정보가 이식되었습니다. (소음/고열 등 {skip_count}개 제외)")
        print(f"[*] 검증이 필요하여 노란색 하이라이트된 행의 수: {len(highlight_rows)}개")
        print(f"[+] 생성된 파일: {output_path}")

    except Exception as e:
        print(f"[!] 엑셀 저장 및 스타일링 중 오류 발생: {e}")

if __name__ == "__main__":
    merge_master_datasets()
