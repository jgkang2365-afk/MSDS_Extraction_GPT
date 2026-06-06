import json
import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# 1. JSON 파일 로드
with open("msds_index.json", "r", encoding="utf-8") as f:
    data = json.load(f)

# 원래 순서를 기록하기 위해 인덱스를 포함한 데이터 리스트 구축
rows = []
for idx, item in enumerate(data):
    row = {"원래순서": idx + 1}
    for k, v in item.items():
        if k == "별칭":
            # 별칭은 JSON 배열 형태 문자열로 저장
            row[k] = json.dumps(v, ensure_ascii=False)
        elif k == "형태학적_필터":
            # 형태학적 필터는 JSON 객체 형태 문자열로 저장 (None인 경우 빈 값)
            if v is not None:
                row[k] = json.dumps(v, ensure_ascii=False)
            else:
                row[k] = ""
        else:
            row[k] = v
    rows.append(row)

df = pd.DataFrame(rows)

# 원래순서 컬럼을 맨 앞으로 이동
cols = ["원래순서"] + [c for c in df.columns if c != "원래순서"]
df = df[cols]

# 엑셀 파일 저장
output_file = "msds_index_edit.xlsx"
try:
    df.to_excel(output_file, index=False)
except PermissionError:
    import sys
    print("❌ 에러: 'msds_index_edit.xlsx' 파일이 엑셀 등의 다른 프로그램에서 열려 있습니다. 엑셀을 완전히 닫은 후 다시 실행해 주십시오.")
    sys.exit(1)

# openpyxl로 엑셀 스타일 적용 (사용 편의성 및 프리미엄 서식 적용)
wb = openpyxl.load_workbook(output_file)
ws = wb.active

# 스타일 정의
header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid") # 차분한 남색 헤더
header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

body_font = Font(name="맑은 고딕", size=10)
body_alignment_center = Alignment(horizontal="center", vertical="center")
body_alignment_left = Alignment(horizontal="left", vertical="center")

thin_side = Side(border_style="thin", color="D3D3D3")
border_all = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

# 헤더 행 스타일 설정
ws.row_dimensions[1].height = 28
for col_idx in range(1, len(cols) + 1):
    cell = ws.cell(row=1, column=col_idx)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_alignment
    cell.border = border_all

# 데이터 행 스타일 설정 및 정렬 적용
for row_idx in range(2, len(df) + 2):
    ws.row_dimensions[row_idx].height = 20
    for col_idx in range(1, len(cols) + 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        cell.font = body_font
        cell.border = border_all
        
        # 특정 열은 가운데 정렬, 나머지는 왼쪽 정렬
        col_name = cols[col_idx - 1]
        if col_name == "CAS No.":
            cell.number_format = '@'
            cell.data_type = 's'
            if cell.value is not None:
                cell.value = str(cell.value)

        if col_name in ["원래순서", "정렬코드", "CAS No.", "적정유속(L/min)", "유속(L/min)", "최소공기량(L)", "최대공기량(L)", "시료수", "허가대상", "특별관리", "측정", "특검", "발암성", "생식세포", "생식독성", "LOD"]:
            cell.alignment = body_alignment_center
        else:
            cell.alignment = body_alignment_left

# 열 너비 자동 조정
for col in ws.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        if cell.value is not None:
            val_str = str(cell.value)
            length = sum(2 if ord(char) > 127 else 1 for char in val_str)
            if length > max_len:
                max_len = length
    # 최소 10, 최대 40 너비 설정
    ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 40)

# 틀 고정 (E열 기준 고정: A~D열 고정 및 1행 고정)
ws.freeze_panes = "E2"

wb.save(output_file)
print(f"성공적으로 {output_file}이 생성되었습니다.")
